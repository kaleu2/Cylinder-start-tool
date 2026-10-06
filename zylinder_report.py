"""Report-Erzeugung (xlsx) fuer das Zylinderabflug-Auswertetool."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    RICH_AVAILABLE = True
except ImportError:  # aeltere openpyxl-Version -> einfacher Text
    RICH_AVAILABLE = False

from i18n import tr
from zylinder_core import (
    seconds_to_hms, PilotResult, SEV_INFO, SEV_WARN, SEV_PENALTY, method_label,
)

HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
FILL_BY_SEV = {
    SEV_PENALTY: PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),  # rot
    SEV_WARN: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),     # gelb
    SEV_INFO: PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"),     # grau
}
FONT_COLOR_BY_SEV = {SEV_PENALTY: "C00000", SEV_WARN: "9C6500", SEV_INFO: "7F7F7F"}
def _sev_label(sev):
    return {SEV_PENALTY: tr("STRAFE/UNGÜLTIG"), SEV_WARN: tr("PRÜFEN"), SEV_INFO: tr("Hinweis")}[sev]


def _fmt_time(t):
    if t is None:
        return ""
    return seconds_to_hms(t)


def _notes_cell_value(notes):
    """Mehrzeilige Besonderheiten, je Zeile in der Farbe ihres Schweregrads."""
    if not notes:
        return ""
    if not RICH_AVAILABLE:
        return "\n".join(f"[{_sev_label(n.sev)}] {n.text}" for n in notes)
    parts = []
    for i, n in enumerate(notes):
        text = f"[{_sev_label(n.sev)}] {n.text}" + ("\n" if i < len(notes) - 1 else "")
        parts.append(TextBlock(
            InlineFont(rFont="Arial", sz=10, color=FONT_COLOR_BY_SEV[n.sev],
                       b=(n.sev == SEV_PENALTY)), text))
    return CellRichText(*parts)


def write_report_xlsx(results: list, out_path: str, task_summary: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = tr("Zylinderabflug")

    ws["A1"] = tr("Zylinderabflug-Auswertung (Annex A 7.4.4)")
    ws["A1"].font = Font(name="Arial", size=14, bold=True)

    row = 3
    for label, value in task_summary.items():
        ws.cell(row=row, column=1, value=label).font = Font(name="Arial", bold=True)
        ws.cell(row=row, column=2, value=value).font = Font(name="Arial")
        row += 1
    row += 1

    # Legende
    ws.cell(row=row, column=1, value=tr("Legende Zeilenfarbe:")).font = Font(name="Arial", bold=True)
    for col, (sev, txt) in enumerate(
            [(SEV_PENALTY, tr("rot = Strafe / ungültig")), (SEV_WARN, tr("gelb = prüfen")),
             (SEV_INFO, tr("grau = nur Hinweis (z. B. Außenlandung, ignorierter PEV)"))], start=2):
        c = ws.cell(row=row, column=col, value=txt)
        c.fill = FILL_BY_SEV[sev]
        c.font = Font(name="Arial")
    row += 2

    headers = [tr(h) for h in [
        "Pilot", "Kennung", "Fluggerät", "Datum (Datei)", "Startmethode", "Startzeit (UTC)",
        "Abstand zum Zylinderrand (m, − = innen)", "Abflughöhe Baro (m)",
        "Groundspeed (km/h)", "Groundspeed-Basis",
        "Ziel erreicht", "Zielzeit (UTC, interpoliert)", "Ankunftshöhe Baro (m, interpoliert)",
        "Höhenverlust Baro (m)", "Strafminuten", "Strafpunkte Speed",
        "Strafpunkte Loss of Height", "Strafpunkte gesamt", "Baro-Kalibrierung",
        "Besonderheiten",
    ]]
    header_row = row
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(name="Arial", bold=True, color="1F4E78")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    row += 1

    thin = Side(style="thin", color="BFBFBF")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for r in results:
        edge = r.start.edge_distance_m if r.start.valid else None
        vals = [
            r.pilot, r.comp_id, r.glider_id, r.flight_date, method_label(r.start.method),
            _fmt_time(r.start.time_s),
            round(edge) if edge is not None else "",
            round(r.start_alt_baro_m) if r.start_alt_baro_m is not None else "",
            round(r.speed_kmh, 1) if r.speed_kmh is not None else "",
            r.speed_detail,
            tr("ja") if r.finish.found else tr("nein"),
            _fmt_time(r.finish.time_s) if r.finish.found else "",
            round(r.finish_alt_baro_m) if r.finish_alt_baro_m is not None else "",
            round(r.loss_of_height_baro_m) if r.loss_of_height_baro_m is not None else "",
            round(r.start.penalty_minutes, 1),
            int(r.speed_penalty_pts),
            int(r.loh_penalty_pts),
            int(r.total_penalty_pts),
            r.baro_cal_text,
            _notes_cell_value(r.notes),
        ]
        sev = r.worst_sev
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=col, value=v)
            if not (RICH_AVAILABLE and isinstance(v, CellRichText)):
                c.font = Font(name="Arial")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = border
            if sev in FILL_BY_SEV:
                c.fill = FILL_BY_SEV[sev]
        if r.date_mismatch:
            ws.cell(row=row, column=4).font = Font(name="Arial", bold=True, color="C00000")
        row += 1

    widths = [18, 8, 10, 12, 11, 11, 14, 11, 11, 14, 8, 13, 14, 11, 10, 10, 12, 11, 24, 80]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[header_row].height = 48

    ws.freeze_panes = ws.cell(row=header_row + 1, column=2).coordinate

    wb.save(out_path)

